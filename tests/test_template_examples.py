"""Exercise the shipped component examples without model or network calls."""
import json
import os
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch
import nbformat
from openpyxl import load_workbook
from zemi import env
from zemi.component import ZemiComponent
from zemi.params import ParamSpace
from zemi.prompting import load_prompts
import tomllib


TEMPLATE = Path(__file__).resolve().parents[2] / 'zemi_component_template'


class TemplateExamplesTests(unittest.TestCase):
    def test_validation_targets_and_independent_parameter_packages(self):
        data = json.loads((TEMPLATE / 'data/validation/validation.json').read_text(encoding='utf-8'))
        self.assertEqual(len(data['items']), 4)
        self.assertEqual([i['ground_truth'] for i in data['items']],
            [['A1:C4'], ['B3:D6'], ['A1:B4', 'D1:E4'], []])
        for item in data['items']:
            file = TEMPLATE / item['input']['workbook_path'].removeprefix('@comp/')
            book = load_workbook(file, data_only=True)
            try:
                self.assertEqual(book.sheetnames, ['Sheet1'])
                sheet = book['Sheet1']
                self.assertTrue(any(cell.value is not None for row in sheet for cell in row))
                for ref in item['ground_truth']:
                    self.assertTrue(all(cell.value is not None for row in sheet[ref] for cell in row))
            finally:
                book.close()
        for folder, count in [('example', 1), ('optimizer_example', 4)]:
            config = tomllib.loads((TEMPLATE / folder / 'params/params.toml').read_text(encoding='utf-8'))
            self.assertEqual(len(ParamSpace(config=config['modules'][0]['params']).grid()), count)
            self.assertEqual('optimizer' in config['modules'][0], folder == 'optimizer_example')
            job = (TEMPLATE / folder / 'job.py').read_text(encoding='utf-8')
            self.assertNotIn('review', job)
        self.assertFalse((TEMPLATE / 'tests').exists())
        self.assertTrue((TEMPLATE / '00_init.py').is_file())

    def test_actual_notebooks_fixed_run_and_optimizer_with_automatic_review(self):
        env.path.tmp.mkdir(parents=True, exist_ok=True)
        cwd = Path.cwd()
        runid = env.path.comp._runid
        with tempfile.TemporaryDirectory(dir=env.path.tmp) as directory:
            root = Path(directory)
            for folder in ('example', 'optimizer_example', 'data'):
                shutil.copytree(TEMPLATE / folder, root / folder)
            shutil.copytree(Path(__file__).resolve().parents[1] / 'zemi', root / 'zemi',
                ignore=shutil.ignore_patterns('.git', '__pycache__'))
            (root / '.zemicomp').touch()
            try:
                os.chdir(root)
                for folder in ('example', 'optimizer_example'):
                    path = root / folder / 'playbook.ipynb'
                    notebook = nbformat.read(path, as_version=4)
                    for cell in notebook.cells:
                        if cell.id == 'select-model':
                            cell.source = '''from types import SimpleNamespace
def fake_completion(**kwargs):
    assert kwargs['messages'][0]['content'] == prompt
    assert 'ground_truth' not in dataset_input
    assert '{{item}}' not in prompt
    targets = {'single_table.xlsx': ['A1:C4'], 'offset_table.xlsx': ['B3:D6'],
               'two_tables.xlsx': ['A1:B4', 'D1:E4'], 'no_tables.xlsx': []}
    import json
    return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(
        {'ranges': targets[dataset_input['workbook_path'].rsplit('/', 1)[1]]})))],
        usage=SimpleNamespace(prompt_tokens=123))
assistant = SimpleNamespace(clients=SimpleNamespace(model='fake'))
client = SimpleNamespace(base_url='http://stub/v1', chat=SimpleNamespace(completions=SimpleNamespace(create=fake_completion)))
'''
                        if cell.id == 'detect-tables':
                            cell.source = cell.source.replace('from urllib.request import Request, urlopen',
                                'from urllib.request import Request\nfrom io import BytesIO\ndef urlopen(*args, **kwargs):\n    return BytesIO(b\'{"tokens":[1,2,3]}\')')
                    nbformat.write(notebook, path)
                    env.path.comp._runid = None
                    component = ZemiComponent(f'@comp/{folder}/params/params.toml')
                    with patch('sys.argv', [str(root / folder / 'job.py')]), patch('zemi.arsenal.ArsenalSession'), patch('zemi.arsenal.begin'), patch('zemi.arsenal.end'):
                        try:
                            component.run()
                        finally:
                            component.close()
                    self.assertEqual(component.report.data['status'], 'succeeded')
                    self.assertEqual(len(component.report.data['trials']), 1 if folder == 'example' else 16)
                    self.assertFalse(list(component.run_directory.glob('**/*.html')))
                    if folder == 'optimizer_example':
                        parent = component.report.data['job_trial']['playbook_trials'][0]
                        self.assertEqual([s['sample_trial_id'] for s in parent['samples']],
                            ['cells_basic-001', 'cells_examples-001', 'cells_compact_basic-001', 'cells_compact_examples-001'])
                        self.assertTrue(all(s['score'] == 1 for s in parent['samples']))
                        review = json.loads((component.run_directory / 'table-detection.review.json').read_text(encoding='utf-8'))
                        self.assertEqual(review['entrypoint'], 'optimizer_example/job.py')
                        self.assertEqual(len(review['prompts']), 4)
                        self.assertIn('data/validation/validation.json', review['sources'])
                        self.assertIn('optimizer_example/params/encoder.py', review['sources'])
                        self.assertEqual(len(review['data_files']), 4)
                        report = (component.run_directory / parent['dataset_report']).read_text(encoding='utf-8')
                        self.assertGreaterEqual(report.count('✅'), 16)
            finally:
                os.chdir(cwd)
                env.path.comp._runid = runid
