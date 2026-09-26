import copy
import hashlib
import json
import os
import shutil
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from zemi import env
from zemi.component import Module, Playbook, ZemiComponent
from zemi.dataset import RunContext, TableDetectionTrialDataset, TrialDataset, resolve_adapter, table_dataset, table_evaluator
from zemi.params import ModuleOptimizer, ParamSpace, PlaybookOptimizer, SampleTrialResult, run_playbook_trial


class DatasetTests(unittest.TestCase):
    def setUp(self):
        env.path.tmp.mkdir(parents=True, exist_ok=True)
        self.root = env.path.tmp / f"dataset-test-{os.getpid()}-{id(self)}"
        self.root.mkdir()
        (self.root / '.zemicomp').touch()
        self.cwd = Path.cwd()
        os.chdir(self.root)
        env.path.comp._runid = None
        book = Workbook()
        book.active.title = 'Данные'
        book.active.append(['City', 'Amount'])
        book.active.append(['A', 10])
        book.create_sheet('Empty')
        book.save('book.xlsx')
        book.close()
        self.data = {'items': [
            {'id': 's', 'description': 'Заголовок', 'input': {'workbook_path': '@comp/book.xlsx', 'worksheet_name': 'Данные'}, 'ground_truth': ['A1:B2'], 'tags': ['header']},
            {'id': 'e', 'input': {'workbook_path': '@comp/book.xlsx', 'worksheet_name': 'Empty'}, 'ground_truth': [], 'tags': ['empty']},
        ]}
        self.write()

    def tearDown(self):
        os.chdir(self.cwd)
        env.path.comp._runid = None
        shutil.rmtree(self.root)

    def write(self):
        Path('data.json').write_text(json.dumps(self.data, ensure_ascii=False), encoding='utf-8')

    def load(self):
        return table_dataset(path='@comp/data.json', params={})

    def test_light_items_and_single_book_preflight(self):
        items = self.load()
        self.assertEqual(items[0]['ground_truth'], ['A1:B2'])
        self.assertEqual(items[1]['ground_truth'], [])
        json.dumps(items)
        self.assertEqual(set(items[0]['input']), {'workbook_path', 'worksheet_name'})

    def test_rejects_all_invalid_contracts(self):
        original = copy.deepcopy(self.data)
        changes = [('id', ''), ('ground_truth', ['B2:A1']), ('tags', 'bad')]
        for key, value in changes:
            with self.subTest(key=key, value=value):
                self.data = copy.deepcopy(original); self.data['items'][0][key] = value
                self.write()
                with self.assertRaises((ValueError, FileNotFoundError)):
                    self.load()

    def test_flat_dataset_unicode_round_trip(self):
        dataset = TrialDataset.load('@comp/data.json')
        self.assertEqual(dataset.items, self.data['items'])
        self.assertIn('Данные', Path('data.json').read_text(encoding='utf-8'))
        self.assertNotIn('\\u0414', Path('data.json').read_text(encoding='utf-8'))

    def test_configured_table_detection_dataset_loads_on_instance(self):
        dataset = TableDetectionTrialDataset(config={'path': '@comp/data.json'})
        self.assertEqual(dataset.items, [])
        dataset.load()
        self.assertEqual(dataset.items, self.data['items'])

    def test_exact_one_to_one_micro_metrics_errors_and_negatives(self):
        items = self.load()
        trial = SampleTrialResult(ParamSpace(config={}).start, [
            {'item': items[0], 'prediction': {'ranges': ['A1:B2', 'A1:B2', 'A1:B3']}},
            {'item': items[1], 'prediction': {'ranges': []}},
            {'item': items[0], 'prediction': None, 'error': 'model failed'},
            {'item': items[1], 'prediction': {'ranges': ['nonsense']}},
        ], {})
        metrics, feedback = table_evaluator(trial, params={})
        self.assertEqual((metrics['tp'], metrics['fp'], metrics['fn']), (1, 3, 1))
        self.assertAlmostEqual(metrics['f1'], 1/3)
        self.assertEqual((metrics['errors'], metrics['reviewed_empty'], metrics['correct_empty']), (2, 2, 1))
        self.assertEqual(len(feedback['items']), 4)
        self.assertEqual(feedback['tags']['header']['fn'], 1)

    def test_adapter_resolver_restricts_local_paths(self):
        Path('adapter.py').write_text('def evaluate(trial, *, params):\n    return {"score": 3}\n')
        self.assertEqual(resolve_adapter('evaluator', '@comp/adapter.py:evaluate')(None, params={}), {'score': 3})
        for name in ('os:system', '@inst/adapter.py:evaluate', '@comp/../adapter.py:evaluate', '@comp/adapter.py:missing'):
            with self.subTest(name=name), self.assertRaises((ValueError, FileNotFoundError)):
                resolve_adapter('evaluator', name)

    def test_cache_closes_before_opening_next_book(self):
        context = RunContext()
        with patch('openpyxl.load_workbook', wraps=load_workbook) as load:
            first = context.workbook(str(self.root / 'book.xlsx'))
            self.assertIs(first, context.workbook(str(self.root / 'book.xlsx')))
            with patch.object(first, 'close', wraps=first.close) as close:
                context.workbook('book.xlsx')
                close.assert_called_once()
            context.close()
        self.assertEqual(load.call_count, 2)

    def component(self):
        Path('one.ipynb').write_text(json.dumps({'cells': [], 'metadata': {}, 'nbformat': 4, 'nbformat_minor': 5}))
        Path('params').mkdir()
        Path('params/test.toml').write_text('''
[system]
version = "0.6"
[component]
[[arsenals]]
id = "local"
config_path = "@comp/arsenal.toml"
lifecycle = "job"
[[modules]]
id = "detect"
kind = "playbook"
path = "one.ipynb"
arsenal = "local"
[modules.params]
x = { values = [0, 1], start = 0 }
[modules.optimizer]
mode = "optimize"
strategy = "block_coordinate"
max_trials = 2
blocks = [["x"]]
[modules.optimizer.sample_trial]
type = "@comp/zemi/sample_trial.py:TableDetectionSampleTrial"
[modules.optimizer.trial_dataset]
path = "@comp/data.json"
''', encoding='utf-8')
        return ZemiComponent('@comp/params/test.toml')

    def test_component_failfast_before_arsenal_and_notebook(self):
        component = self.component()
        self.data['items'][1]['ground_truth'] = ['bad']
        self.write()
        with patch('zemi.arsenal.begin') as begin, patch('zemi.component.Playbook.run') as run:
            with self.assertRaisesRegex(ValueError, 'Invalid exact range'):
                component.run()
            begin.assert_not_called()
            run.assert_not_called()
        component.close()

    def test_component_lifecycle_report_and_no_ground_truth_in_notebook(self):
        component = self.component()
        (self.root / 'zemi').mkdir()
        (self.root / 'job.py').write_text('# test entrypoint', encoding='utf-8')
        component.reporting.configure_review('detect', entrypoint='@comp/job.py',
            prompts={'test': '### Input\n{{worksheet_text}}\n### Output\n{"ranges":[]}'},
            sources=['@comp/data.json'])
        captured = []
        def notebook(playbook):
            captured.append(copy.deepcopy(playbook.params))
            entry = component.report.start_trial(playbook)
            name = playbook.params['dataset_input']['worksheet_name']
            if playbook.params['x'] == 0 and name == 'Данные':
                error = RuntimeError('model unavailable')
                component.report.fail_playbook(entry, error)
                raise error
            entry['output_params'] = {'ranges': ['A1:B2'] if name == 'Данные' else []}
            component.report.finish_playbook(entry)
        with patch('zemi.arsenal.ArsenalSession'), patch('zemi.arsenal.begin') as begin, patch('zemi.arsenal.end') as end, patch('zemi.component.Playbook.run', notebook):
            component.run()
            begin.assert_called_once()
            end.assert_called_once()
        component.close()
        review = (component.run_directory / 'detect.review.md').read_text(encoding='utf-8')
        self.assertIn('| Status | failed |', review)
        self.assertIn('| Successful runs / Total | 3 / 4 |', review)
        self.assertIn('detect.review.json', review)
        self.assertTrue((component.run_directory / 'detect.review.json').is_file())
        for obsolete in ('main.md', 'report.md'):
            self.assertFalse((component.run_directory / obsolete).exists())
        self.assertFalse((component.run_directory / 'sample_trials').exists())
        self.assertFalse(list((component.run_directory / 'notebooks').glob('*.report.md')))
        report = json.loads(component.report.path.read_text(encoding='utf-8'))
        parent = report['job_trial']['playbook_trials'][0]
        self.assertEqual(parent['optimizer']['blocks'], [['x']])
        module_path = component.run_directory / component.reporting.writer.ref('module', 'detect').path
        self.assertIn('## Module Optimization Progress', module_path.read_text(encoding='utf-8'))
        self.assertIn('[["x"]]', module_path.read_text(encoding='utf-8'))
        self.assertEqual(len(captured), 4)
        self.assertTrue(all(set(p['dataset_input']) == {'workbook_path', 'worksheet_name'} for p in captured))
        self.assertNotIn('reference', json.dumps(captured))
        self.assertEqual(len(parent['samples']), 2)
        self.assertTrue(all(len(s['runs']) == 2 for s in parent['samples']))
        self.assertEqual(parent['best_sample'], 'detect-sample-0002')
        self.assertEqual(parent['samples'][0]['metrics']['fn'], 1)
        self.assertEqual(parent['samples'][1]['metrics']['f1'], 1)
        self.assertIn('`optimize`', (component.run_directory / 'index.md').read_text(encoding='utf-8'))
        sample_report = component.run_directory / parent['samples'][0]['report']
        dataset_report = component.run_directory / parent['dataset_report']
        progress_report = component.run_directory / component.reporting.writer.ref('module', 'detect').path
        for target in (sample_report, dataset_report, progress_report):
            self.assertTrue(target.is_file())
        self.assertIn('Ground truth', sample_report.read_text(encoding='utf-8'))
        self.assertIn('model unavailable', sample_report.read_text(encoding='utf-8'))
        self.assertIn('Данные', sample_report.read_text(encoding='utf-8'))
        progress = progress_report.read_text(encoding='utf-8')
        self.assertIn(f'({dataset_report.name})', progress)
        self.assertIn('## Module Optimization Progress', progress)
        self.assertIn('## Selected Sample', progress)
        self.assertIn('Worksheets detected', dataset_report.read_text(encoding='utf-8'))
        raw_report = component.report.path.read_text(encoding='utf-8')
        self.assertIn('"worksheet_name": "Данные"', raw_report)
        self.assertNotIn('\\u0414', raw_report)

    def test_real_papermill_dataset_smoke_without_model(self):
        import nbformat
        component = self.component()
        library_parent = str(Path(__file__).resolve().parents[1])
        notebook = nbformat.v4.new_notebook()
        notebook.metadata.kernelspec = {'display_name': 'Python 3', 'language': 'python', 'name': 'python3'}
        notebook.cells = [
            nbformat.v4.new_code_cell('x = 0\ndataset_input = {}\narsenal_config_path = ""\narsenal_start_and_stop_at_job_level = True', metadata={'tags': ['parameters']}),
            nbformat.v4.new_code_cell(f'import sys, os\nsys.path.insert(0, {library_parent!r})\nfrom zemi.playbook import output_params\nassert "previous_run" not in globals()\nprevious_run = True\nassert set(dataset_input) == {{"workbook_path", "worksheet_name"}}\noutput_params({{"ranges": [], "pid": os.getpid()}})'),
        ]
        nbformat.write(notebook, 'one.ipynb')
        # Freeze a valid sample using the public holdout API.
        component.close()
        env.path.comp._runid = None
        component = ZemiComponent('@comp/params/test.toml', sample_overrides={'detect': {'x': 0}})
        with patch('zemi.arsenal.ArsenalSession'), patch('zemi.arsenal.begin'), patch('zemi.arsenal.end'):
            try:
                component.run()
            finally:
                component.close()
        parent = component.report.data['job_trial']['playbook_trials'][0]
        self.assertEqual(len(parent['samples']), 1)
        self.assertEqual(len(parent['samples'][0]['runs']), 2)
        self.assertEqual(parent['samples'][0]['metrics']['fn'], 1)
        self.assertEqual(parent['samples'][0]['metrics']['correct_empty'], 1, parent['samples'][0]['runs'])
        for run in parent['samples'][0]['runs']:
            self.assertTrue((component.run_directory / run['artifacts']['output_notebook']).is_file())
            self.assertEqual(run['prediction']['ranges'], [])
        self.assertEqual(len({r['prediction']['pid'] for r in parent['samples'][0]['runs']}), 1)
        self.assertFalse(component._module_kernels)

    def test_shared_kernel_restarts_after_error_and_does_not_leak_variables(self):
        import nbformat
        component = self.component()
        library_parent = str(Path(__file__).resolve().parents[1])
        nb = nbformat.v4.new_notebook(metadata={'kernelspec': {'display_name': 'Python 3', 'language': 'python', 'name': 'python3'}})
        nb.cells = [
            nbformat.v4.new_code_cell('x = 0\ndataset_input = {}\narsenal_config_path = ""\narsenal_start_and_stop_at_job_level = True', metadata={'tags': ['parameters']}),
            nbformat.v4.new_code_cell(f'import sys, os\nsys.path.insert(0, {library_parent!r})\nfrom zemi.playbook import output_params\nassert "previous_run" not in globals()\nprevious_run = True\noutput_params({{"ranges": ["A1:B2"] if dataset_input["worksheet_name"] == "Данные" else [], "pid": os.getpid()}})\nif x == 0 and dataset_input["worksheet_name"] == "Данные":\n    raise RuntimeError("intentional failure")'),
        ]
        nbformat.write(nb, 'one.ipynb')
        with patch('zemi.arsenal.ArsenalSession'), patch('zemi.arsenal.begin'), patch('zemi.arsenal.end'):
            try:
                component.run()
            finally:
                component.close()
        trials = component.report.data['trials']
        self.assertEqual([t['status'] for t in trials], ['failed', 'succeeded', 'succeeded', 'succeeded'])
        self.assertNotEqual(trials[0]['output_params']['pid'], trials[1]['output_params']['pid'])
        self.assertEqual(len({t['output_params']['pid'] for t in trials[1:]}), 1)
        self.assertFalse(component._module_kernels)
        self.assertTrue(all((component.run_directory / t['output_notebook']).is_file() for t in trials))

    def test_reuse_kernel_can_be_disabled(self):
        import nbformat
        component = self.component()
        config = self.root / 'params/test.toml'
        config.write_text(config.read_text(encoding='utf-8').replace('max_trials = 2', 'max_trials = 2\nreuse_kernel = false'), encoding='utf-8')
        component.close()
        env.path.comp._runid = None
        component = ZemiComponent('@comp/params/test.toml', sample_overrides={'detect': {'x': 0}})
        library_parent = str(Path(__file__).resolve().parents[1])
        nb = nbformat.v4.new_notebook(metadata={'kernelspec': {'display_name': 'Python 3', 'language': 'python', 'name': 'python3'}})
        nb.cells = [
            nbformat.v4.new_code_cell('x = 0\ndataset_input = {}\narsenal_config_path = ""\narsenal_start_and_stop_at_job_level = True', metadata={'tags': ['parameters']}),
            nbformat.v4.new_code_cell(f'import sys, os\nsys.path.insert(0, {library_parent!r})\nfrom zemi.playbook import output_params\noutput_params({{"ranges": [], "pid": os.getpid()}})'),
        ]
        nbformat.write(nb, 'one.ipynb')
        with patch('zemi.arsenal.ArsenalSession'), patch('zemi.arsenal.begin'), patch('zemi.arsenal.end'):
            try:
                component.run()
            finally:
                component.close()
        self.assertEqual(len({t['output_params']['pid'] for t in component.report.data['trials']}), 2)
        self.assertFalse(component._module_kernels)

    def test_component_can_replay_best_sample_from_report(self):
        component = self.component()
        report_path = self.root / 'validation-report.json'
        report_path.write_text(json.dumps({'job_trial': {'playbook_trials': [{
            'playbook_id': 'detect', 'status': 'succeeded', 'best_sample': 'detect-sample-0002',
            'samples': [
                {'sample_trial_id': 'detect-sample-0001', 'params': {'x': 0}},
                {'sample_trial_id': 'detect-sample-0002', 'params': {'x': 1}},
            ],
        }]}}), encoding='utf-8')
        component.close()
        env.path.comp._runid = None
        replay = ZemiComponent.from_best_report('@comp/params/test.toml', report_path)
        try:
            self.assertEqual(replay.playbooks[0].params['x'], 1)
            self.assertEqual(replay.modules[0].optimizer_config['max_trials'], 1)
        finally:
            replay.close()

    def test_custom_sample_trial_is_the_single_extension_point(self):
        Path('one.ipynb').write_text(json.dumps({'cells': [], 'metadata': {}, 'nbformat': 4, 'nbformat_minor': 5}))
        Path('custom_trial.py').write_text('''
from zemi.sample_trial import SampleTrial

class CustomTrial(SampleTrial):
    def evaluate(self, runs):
        assert self.dataset.items[0]["ground_truth"] == []
        quality = runs[0]["prediction"]["quality"]
        return {"quality": quality, "diagnostic_count": len(runs)}, quality, {"kind": "custom"}

    def render_report(self, runs, metrics, score, feedback):
        return "### Custom SampleTrial report\\n\\nDomain-owned content."
''', encoding='utf-8')
        Path('params').mkdir(exist_ok=True)
        Path('params/custom.toml').write_text('''
[system]
version = "0.6"
[component]
[[modules]]
id = "custom"
kind = "playbook"
path = "one.ipynb"
[modules.params]
x = { values = [0, 1], start = 0 }
[modules.optimizer]
mode = "optimize"
strategy = "grid"
[modules.optimizer.sample_trial]
type = "@comp/custom_trial.py:CustomTrial"
[modules.optimizer.trial_dataset]
path = "@comp/data.json"
''', encoding='utf-8')
        Path('data.json').write_text(json.dumps({"items": [{"id": "only", "input": {"value": 1}, "ground_truth": []}]}), encoding='utf-8')
        component = ZemiComponent('@comp/params/custom.toml')
        def notebook(playbook):
            entry = component.report.start_trial(playbook)
            entry['output_params'] = {'quality': playbook.params['x']}
            component.report.finish_playbook(entry)
        with patch('zemi.component.Playbook.run', notebook):
            component.run()
        component.close()
        parent = component.report.data['job_trial']['playbook_trials'][0]
        self.assertEqual([sample['score'] for sample in parent['samples']], [0.0, 1.0])
        self.assertEqual(parent['samples'][1]['metrics'], {'quality': 1.0, 'diagnostic_count': 1.0})
        self.assertEqual(parent['samples'][1]['feedback'], {'kind': 'custom'})
        self.assertEqual(parent['best_params'], {'x': 1})
        detailed = component.run_directory / parent['samples'][1]['report']
        self.assertTrue(detailed.is_file())
        self.assertIn('Domain-owned content', detailed.read_text(encoding='utf-8'))

    def test_component_module_hierarchy_and_start_only_unicode(self):
        component = self.component()
        self.assertEqual(component.modules, component.playbooks)
        self.assertIsInstance(component.modules[0], Module)
        self.assertIsInstance(component.modules[0], Playbook)
        self.assertEqual(component.modules[0].kind, 'playbook')
        component.close()

    def test_start_only_via_select_runs_one_full_sample_trial(self):
        initial = self.component()
        initial.close()
        env.path.comp._runid = None
        config_path = Path('params/test.toml')
        text = config_path.read_text(encoding='utf-8').replace(
            'mode = "optimize"', 'mode = { select = ["optimize", "start_only"] }'
        )
        config_path.write_text(text, encoding='utf-8')
        with patch('builtins.input', return_value='2'):
            component = ZemiComponent('@comp/params/test.toml')

        def notebook(playbook):
            entry = component.report.start_trial(playbook)
            entry['output_params'] = {'ranges': []}
            component.report.finish_playbook(entry)

        with patch('zemi.arsenal.ArsenalSession'), patch('zemi.arsenal.begin'), patch('zemi.arsenal.end'), patch('zemi.component.Playbook.run', notebook):
            component.run()
        component.close()
        trial = component.report.data['job_trial']['playbook_trials'][0]
        self.assertEqual(trial['optimizer']['mode'], 'start_only')
        self.assertEqual(len(trial['samples']), 1)
        self.assertEqual(len(trial['samples'][0]['runs']), 2)
        self.assertEqual(trial['best_params'], {'x': 0})
        module_doc = (component.run_directory / component.reporting.writer.ref('module', 'detect').path).read_text(encoding='utf-8')
        self.assertIn('No parameter search was performed.', module_doc)
        self.assertIn('`start_only`', (component.run_directory / 'index.md').read_text(encoding='utf-8'))

    def test_best_report_must_resolve_a_successful_sample(self):
        self.component().close()
        for payload in (
            {},
            {'job_trial': {'playbook_trials': [{'playbook_id': 'detect', 'status': 'failed'}]}},
            {'job_trial': {'playbook_trials': [{'playbook_id': 'detect', 'status': 'succeeded',
                                               'best_sample': 'missing', 'samples': []}]}},
        ):
            with self.subTest(payload=payload):
                report_path = self.root / 'bad-report.json'
                report_path.write_text(json.dumps(payload), encoding='utf-8')
                with self.assertRaisesRegex(ValueError, 'best sample report'):
                    ZemiComponent.from_best_report('@comp/params/test.toml', report_path)


class AdaptiveTests(unittest.TestCase):
    def test_optimizer_report_truncates_only_long_param_samples(self):
        optimizer = ModuleOptimizer(config={'strategy': 'grid'}, param_space=ParamSpace(config={}))
        short = SampleTrialResult(ParamSpace(config={'x': 1}).start, [], {'f1': 1}, score=1, report='short.md')
        long = SampleTrialResult(ParamSpace(config={'first': 'x' * 50, 'second': 'y' * 50}).start, [], {'f1': .5}, score=.5, report='long.md')
        report = optimizer.render_report(history=[short, long], best_param_sample=short.sample)
        self.assertIn('x=1', report)
        self.assertNotIn('x=1 / …', report)
        self.assertIn(' / …', report)
        self.assertIn('[Best](short.md)', report)
        self.assertIn('[Details](long.md)', report)

    def test_optimizer_next_and_best_use_history_score(self):
        space = ParamSpace(config={'x': {'values': [0, 1], 'start': 0}})
        optimizer = PlaybookOptimizer(config={'strategy': 'grid'}, param_space=space)
        first = optimizer.next_param_sample([])
        history = [SampleTrialResult(first, [], {'loss': 3}, score=3)]
        second = optimizer.next_param_sample(history)
        history.append(SampleTrialResult(second, [], {'loss': 1, 'other': 9}, score=4))
        self.assertEqual(optimizer.best_param_sample(history).values, {'x': 1})
        self.assertIsNone(optimizer.next_param_sample(history))

    def test_coordinate_follows_improving_anchor(self):
        for strategy, extra in [('coordinate', {}), ('block_coordinate', {'blocks': [['x'], ['y']]})]:
            with self.subTest(strategy=strategy):
                space = ParamSpace(config={key: {'values': [0, 1], 'start': 0} for key in ('x', 'y')})
                optimizer = PlaybookOptimizer(
                    config={'strategy': strategy, 'max_samples': 4, **extra}, param_space=space
                )
                result = run_playbook_trial(optimizer=optimizer, dataset=[None], run=lambda s, i: None,
                    evaluator=lambda s, r: ({'quality': 2*s.values['x']+s.values['y']}, 2*s.values['x']+s.values['y']))
                self.assertEqual(optimizer.best_param_sample(result.history).values, {'x': 1, 'y': 1})

    def test_block_coordinate_can_improve_jointly_with_unlisted_singletons(self):
        space = ParamSpace(config={
            'x': {'values': [0, 1], 'start': 0},
            'y': {'values': [0, 1], 'start': 0},
            'z': {'values': [0, 1], 'start': 0},
        })
        optimizer = PlaybookOptimizer(
            config={'strategy': 'block_coordinate', 'max_samples': 6, 'blocks': [['x', 'y']]},
            param_space=space,
        )
        result = run_playbook_trial(
            optimizer=optimizer, dataset=[None], run=lambda sample, item: None,
            evaluator=lambda sample, runs: ({'quality': 10 * (sample.values['x'] == sample.values['y'] == 1) + sample.values['z']},
                10 * (sample.values['x'] == sample.values['y'] == 1) + sample.values['z']),
        )
        self.assertEqual(optimizer.best_param_sample(result.history).values, {'x': 1, 'y': 1, 'z': 1})

    def test_block_coordinate_order_restart_uniqueness_and_exhaustion(self):
        space = ParamSpace(config={
            'x': {'values': [0, 1], 'start': 0},
            'y': {'values': [0, 1], 'start': 0},
        })
        optimizer = PlaybookOptimizer(
            config={'strategy': 'block_coordinate', 'max_samples': 10, 'blocks': [['x'], ['y']]},
            param_space=space,
        )
        scores = {(0, 0): 0, (1, 0): 0, (0, 1): 10, (1, 1): 11}
        result = run_playbook_trial(
            optimizer=optimizer, dataset=[None], run=lambda sample, item: None,
            evaluator=lambda sample, runs: ({'quality': scores[(sample.values['x'], sample.values['y'])]},
                                             scores[(sample.values['x'], sample.values['y'])]),
        )
        values = [item.sample.values for item in result.history]
        self.assertEqual(values, [
            {'x': 0, 'y': 0}, {'x': 1, 'y': 0},
            {'x': 0, 'y': 1}, {'x': 1, 'y': 1},
        ])
        self.assertEqual(len({item.sample.key() for item in result.history}), len(result.history))
        self.assertEqual(len(result.history), 4)  # finite space exhausted before max_samples

        limited = PlaybookOptimizer(
            config={'strategy': 'block_coordinate', 'max_samples': 3, 'blocks': [['x'], ['y']]},
            param_space=space,
        )
        limited_result = run_playbook_trial(
            optimizer=limited, dataset=[None], run=lambda sample, item: None,
            evaluator=lambda sample, runs: ({'quality': 0}, 0),
        )
        self.assertEqual(len(limited_result.history), 3)

    def test_block_coordinate_respects_declared_block_order(self):
        space = ParamSpace(config={
            'x': {'values': [0, 1], 'start': 0},
            'y': {'values': [0, 1], 'start': 0},
        })
        optimizer = PlaybookOptimizer(
            config={'strategy': 'block_coordinate', 'max_samples': 2, 'blocks': [['y'], ['x']]},
            param_space=space,
        )
        result = run_playbook_trial(
            optimizer=optimizer, dataset=[None], run=lambda sample, item: None,
            evaluator=lambda sample, runs: ({'quality': 0}, 0),
        )
        self.assertEqual(result.history[1].sample.values, {'x': 0, 'y': 1})

    def test_invalid_score_is_failed_observation_and_run_failure_retained(self):
        optimizer = PlaybookOptimizer(
            config={'strategy': 'grid'},
            param_space=ParamSpace(config={'x': {'values': [0, 1], 'start': 0}}),
        )
        def run(sample, item):
            raise RuntimeError('failure')
        result = run_playbook_trial(optimizer=optimizer, dataset=[1, 2], run=run,
            evaluator=lambda s, r: ({'quality': 0}, float('nan')))
        self.assertEqual(len(result.history), 2)
        self.assertTrue(all(s.error and len(s.runs) == 2 for s in result.history))
        self.assertIsNone(result.best('score', 'maximize'))
